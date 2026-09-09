from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = "CDPS_Pre_Class_Notice.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "DFKai-SB")

def font(run, size=11, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "DFKai-SB")

def paragraph(text="", bold=False, size=11, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    font(r, size, bold)
    return p

def section_heading(text):
    p = paragraph(text, bold=True, size=12, space_after=2)
    p.paragraph_format.space_before = Pt(8)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    font(r)
    return p

def link(text):
    p = paragraph(text, space_after=4)
    for r in p.runs:
        r.font.color.rgb = RGBColor(5, 99, 193)
        r.font.underline = True
    return p

title = paragraph("CDPS Pre Class Notice", bold=True, size=16, space_after=10)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

paragraph("Dear Students,", space_after=5)
paragraph("Welcome to Co-opetitive Dynamics and Platform Strategy (CDPS)!")
paragraph("Please note the following important information before class:")

section_heading("1. Class Time and Location")
paragraph("Monday, 14:30–17:20 (fgh)")
paragraph("B15, College of Management")
paragraph("Please arrive early or on time.")

section_heading("2. Course Expectations")
paragraph("EMI Course", bold=True)
paragraph("This course is conducted entirely in English. Commitment and active participation are essential to your learning experience.")
paragraph("Mixed Grouping", bold=True)
paragraph("Groups will be arranged by the instructor with students from different departments and years. There will be weekly pre-class group assignments and in-class discussions. Diverse backgrounds and active peer support can enhance the learning experience.")
paragraph("Responsible AI Usage", bold=True)
paragraph("AI tools will be used in this course. Please bring a laptop or tablet to class every week. Students are expected to use AI responsibly, remain human-in-the-loop (HITL), and take final accountability for their own learning process and outcomes.")

section_heading("3. Pre-Class Survey")
paragraph("To facilitate mixed grouping, please complete the Pre-Class Survey by 12:00 noon on Sunday, September 6.")
link("https://forms.gle/zrGR3je4VDUwvXgq9")
paragraph("Please refer to the CEFR English Learning Reference in Section 4 when completing the survey.")

section_heading("4. Digital Learning Resources")
paragraph("Professor Lo’s Bio:", bold=True)
link("https://www.ibs.ncnu.edu.tw/index.php/faculty/tenured-professor/299-smlo")
paragraph("CEFR English Learning Reference:", bold=True)
link("https://lolopodcast.github.io/EMI-Rubrics-CEFR2020/")
paragraph("English Self-Introduction Activity:", bold=True)
paragraph("In Week 1, we will introduce a short English self-introduction activity supported by AI. Please briefly explore the webpage before class; detailed instructions and submission requirements will be explained in class.")
link("https://lolopodcast.github.io/EMI-AI-Self-Introduction/")

section_heading("5. Leave Requests")
paragraph("If you need to request leave, please submit it in advance through the “Request for Leave of Absence & Self-Study Note” activity on Moodle.")
paragraph("Example: 2/25, G9/IBS 3/Tina Lin, Sick Leave")
paragraph("(Chinese or English is accepted.)")
paragraph("If you have any questions, please contact the TAs:")
paragraph("Tina Lin: s113212016@mail1.ncnu.edu.tw")
paragraph("Tina Wang: s112212015@mail1.ncnu.edu.tw")
paragraph("Best regards,")
paragraph("The CDPS Team & Prof. Lo")

doc.add_page_break()
title = paragraph("CDPS 課前重要注意事項", bold=True, size=16, space_after=10)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

paragraph("各位同學大家好：", space_after=5)
paragraph("歡迎來到「競合動態與平台策略（Co-opetitive Dynamics and Platform Strategy, CDPS）」課程！")
paragraph("以下為上課前的重要注意事項：")

section_heading("1. 上課時間與地點")
paragraph("每週一 14:30–17:20（fgh）")
paragraph("管理學院 B15 教室")
paragraph("請提前或準時進入教室就座。")

section_heading("2. 課程期待")
paragraph("全英語授課", bold=True)
paragraph("本課程採全英語授課。積極投入與主動參與，是獲得良好學習成效的重要條件。")
paragraph("混合分組", bold=True)
paragraph("課程將由老師依不同系所、年級等背景進行混合編組。每週皆會有課前小組作業與課堂討論；多元背景與積極的同儕支持，能提升整體學習成效。")
paragraph("負責任的 AI 使用", bold=True)
paragraph("本課程將使用 AI 工具進行學習與實作，請每週攜帶筆電或平板電腦。使用 AI 時，請遵守負責任使用原則，保持 Human-in-the-Loop（HITL），並對自己的學習過程與成果負最終責任。")

section_heading("3. 課前表單")
paragraph("為利進行混合分組，請於 9 月 6 日（日）中午 12:00 前完成課前表單填寫。")
link("https://forms.gle/zrGR3je4VDUwvXgq9")
paragraph("填寫表單時，請參考第 4 點的 CEFR 英文程度參考資源。")

section_heading("4. 數位學習資源")
paragraph("駱世民教授個人簡介：", bold=True)
link("https://www.ibs.ncnu.edu.tw/index.php/faculty/tenured-professor/299-smlo")
paragraph("CEFR 英文程度參考資源：", bold=True)
link("https://lolopodcast.github.io/EMI-Rubrics-CEFR2020/")
paragraph("英文自我介紹活動：", bold=True)
paragraph("第一週將介紹一項 AI 輔助的英文自我介紹活動。請同學於課前先簡單瀏覽網頁；詳細操作方式與繳交規定將於課堂中說明。")
link("https://lolopodcast.github.io/EMI-AI-Self-Introduction/")

section_heading("5. 請假規範")
paragraph("若需請假，請務必預先於 Moodle 系統的「請假與自主學習紀錄（Request for Leave of Absence & Self-Study Note）」登記。")
paragraph("登記範例：2/25, G9/國企三/林庭伃, 病假")
paragraph("（中英文皆可）")
paragraph("如有任何問題，歡迎聯繫助教：")
paragraph("Tina Lin：s113212016@mail1.ncnu.edu.tw")
paragraph("Tina Wang：s112212015@mail1.ncnu.edu.tw")
paragraph("祝 學習順利")
paragraph("CDPS 課程團隊暨駱世民教授")

doc.save(OUT)

html = r'''<!doctype html>
<html lang="en">
<head><meta charset="utf-8"><title>CDPS Pre Class Notice</title></head>
<body style="font-size: 11pt; line-height: 1.5;">
<div class="en" style="font-family: 'Times New Roman';">
<p>Dear Students,</p>
<p>Welcome to Co-opetitive Dynamics and Platform Strategy (CDPS)!</p>
<p>Please note the following important information before class:</p>
<p><strong>1. Class Time and Location</strong><br>Monday, 14:30-17:20 (fgh)<br>B15, College of Management<br>Please arrive early or on time.</p>
<p><strong>2. Course Expectations</strong></p>
<p><strong>EMI Course</strong><br>This course is conducted entirely in English. Commitment and active participation are essential to your learning experience.</p>
<p><strong>Mixed Grouping</strong><br>Groups will be arranged by the instructor with students from different departments and years. There will be weekly pre-class group assignments and in-class discussions. Diverse backgrounds and active peer support can enhance the learning experience.</p>
<p><strong>Responsible AI Usage</strong><br>AI tools will be used in this course. Please bring a laptop or tablet to class every week. Students are expected to use AI responsibly, remain human-in-the-loop (HITL), and take final accountability for their own learning process and outcomes.</p>
<p><strong>3. Pre-Class Survey</strong><br>To facilitate mixed grouping, please complete the Pre-Class Survey by 12:00 noon on Sunday, September 6.</p>
<p><a href="https://forms.gle/zrGR3je4VDUwvXgq9">https://forms.gle/zrGR3je4VDUwvXgq9</a></p>
<p>Please refer to the CEFR English Learning Reference in Section 4 when completing the survey.</p>
<p><strong>4. Digital Learning Resources</strong><br>Professor Lo's Bio:<br><a href="https://www.ibs.ncnu.edu.tw/index.php/faculty/tenured-professor/299-smlo">https://www.ibs.ncnu.edu.tw/index.php/faculty/tenured-professor/299-smlo</a></p>
<p>CEFR English Learning Reference:<br><a href="https://lolopodcast.github.io/EMI-Rubrics-CEFR2020/">https://lolopodcast.github.io/EMI-Rubrics-CEFR2020/</a></p>
<p>English Self-Introduction Activity:<br>In Week 1, we will introduce a short English self-introduction activity supported by AI. Please briefly explore the webpage before class; detailed instructions and submission requirements will be explained in class.</p>
<p><a href="https://lolopodcast.github.io/EMI-AI-Self-Introduction/">https://lolopodcast.github.io/EMI-AI-Self-Introduction/</a></p>
<p><strong>5. Leave Requests</strong><br>If you need to request leave, please submit it in advance through the “Request for Leave of Absence &amp; Self-Study Note” activity on Moodle.</p>
<p>Example: 2/25, G9/IBS 3/Tina Lin, Sick Leave<br>(Chinese or English is accepted.)</p>
<p>If you have any questions, please contact the TAs:</p>
<p>Tina Lin: s113212016@mail1.ncnu.edu.tw<br>Tina Wang: s112212015@mail1.ncnu.edu.tw</p>
<p>Best regards,<br>The CDPS Team &amp; Prof. Lo</p>
</div>
<hr>
<div lang="zh-TW" class="zh" style="font-family: 'DFKai-SB';">
<p>各位同學大家好：</p>
<p>歡迎來到「競合動態與平台策略（Co-opetitive Dynamics and Platform Strategy, CDPS）」課程！</p>
<p>以下為上課前的重要注意事項：</p>
<p><strong>1. 上課時間與地點</strong><br>每週一 14:30-17:20（fgh）<br>管理學院 B15 教室<br>請提前或準時進入教室就座。</p>
<p><strong>2. 課程期待</strong></p>
<p><strong>全英語授課</strong><br>本課程採全英語授課。積極投入與主動參與，是獲得良好學習成效的重要條件。</p>
<p><strong>混合分組</strong><br>課程將由老師依不同系所、年級等背景進行混合編組。每週皆會有課前小組作業與課堂討論；多元背景與積極的同儕支持，能提升整體學習成效。</p>
<p><strong>負責任的 AI 使用</strong><br>本課程將使用 AI 工具進行學習與實作，請每週攜帶筆電或平板電腦。使用 AI 時，請遵守負責任使用原則，保持 Human-in-the-Loop（HITL），並對自己的學習過程與成果負最終責任。</p>
<p><strong>3. 課前表單</strong><br>為利進行混合分組，請於 9 月 6 日（日）中午 12:00 前完成課前表單填寫。</p>
<p><a href="https://forms.gle/zrGR3je4VDUwvXgq9">https://forms.gle/zrGR3je4VDUwvXgq9</a></p>
<p>填寫表單時，請參考第 4 點的 CEFR 英文程度參考資源。</p>
<p><strong>4. 數位學習資源</strong><br>駱世民教授個人簡介：<br><a href="https://www.ibs.ncnu.edu.tw/index.php/faculty/tenured-professor/299-smlo">https://www.ibs.ncnu.edu.tw/index.php/faculty/tenured-professor/299-smlo</a></p>
<p>CEFR 英文程度參考資源：<br><a href="https://lolopodcast.github.io/EMI-Rubrics-CEFR2020/">https://lolopodcast.github.io/EMI-Rubrics-CEFR2020/</a></p>
<p>英文自我介紹活動：<br>第一週將介紹一項 AI 輔助的英文自我介紹活動。請同學於課前先簡單瀏覽網頁；詳細操作方式與繳交規定將於課堂中說明。</p>
<p><a href="https://lolopodcast.github.io/EMI-AI-Self-Introduction/">https://lolopodcast.github.io/EMI-AI-Self-Introduction/</a></p>
<p><strong>5. 請假規範</strong><br>若需請假，請務必預先於 Moodle 系統的「請假與自主學習紀錄（Request for Leave of Absence &amp; Self-Study Note）」登記。</p>
<p>登記範例：2/25, G9/國企三/林庭伃, 病假<br>（中英文皆可）</p>
<p>如有任何問題，歡迎聯繫助教：</p>
<p>Tina Lin：s113212016@mail1.ncnu.edu.tw<br>Tina Wang：s112212015@mail1.ncnu.edu.tw</p>
<p>祝 學習順利<br>CDPS 課程團隊暨駱世民教授</p>
</div>
</body></html>'''
with open("CDPS_Pre_Class_Notice_Moodle.html", "w", encoding="utf-8") as stream:
    stream.write(html)
